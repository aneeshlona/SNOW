import os
import requests
from openpyxl import load_workbook
from dotenv import load_dotenv


# --------------------------------------------------
# Load environment variables
# --------------------------------------------------

load_dotenv()

SNOW_URL = os.getenv("SNOW_URL")
SNOW_USERNAME = os.getenv("SNOW_USERNAME")
SNOW_PASSWORD = os.getenv("SNOW_PASSWORD")


# --------------------------------------------------
# Excel file
# --------------------------------------------------

EXCEL_FILE = "batch_job_incidents.xlsx"


# --------------------------------------------------
# ServiceNow URL
# --------------------------------------------------

url = f"{SNOW_URL}/api/now/table/incident"


# --------------------------------------------------
# Read incidents from Excel
# --------------------------------------------------

workbook = load_workbook(EXCEL_FILE)

worksheet = workbook["Incidents"]

# Get column positions from header row
headers = {}

for cell in worksheet[1]:
    headers[cell.value] = cell.column


sys_id_column = headers.get("Sys ID")
number_column = headers.get("Number")
short_description_column = headers.get("Short Description")


if not sys_id_column:
    raise ValueError("Sys ID column was not found in the Excel file")


incidents = []


for row in range(2, worksheet.max_row + 1):

    sys_id = worksheet.cell(
        row=row,
        column=sys_id_column
    ).value

    number = worksheet.cell(
        row=row,
        column=number_column
    ).value

    short_description = worksheet.cell(
        row=row,
        column=short_description_column
    ).value

    if sys_id:
        incidents.append({
            "sys_id": str(sys_id),
            "number": str(number),
            "short_description": str(short_description)
        })


# --------------------------------------------------
# Display incidents that will be deleted
# --------------------------------------------------

print()
print("=" * 80)
print(f"INCIDENTS FOUND IN EXCEL: {len(incidents)}")
print("=" * 80)

for incident in incidents:

    print(
        f"{incident['number']} | "
        f"{incident['short_description']} | "
        f"{incident['sys_id']}"
    )


# --------------------------------------------------
# Confirmation
# --------------------------------------------------

print()
print("=" * 80)
print("WARNING: These incidents will be permanently deleted")
print("=" * 80)

confirmation = input(
    "\nType DELETE to continue: "
)

if confirmation != "DELETE":

    print("\nDeletion cancelled.")
    exit()


# --------------------------------------------------
# Delete incidents
# --------------------------------------------------

deleted = 0
failed = 0 


for incident in incidents:

    sys_id = incident["sys_id"]

    delete_url = f"{url}/{sys_id}"

    response = requests.delete(
        delete_url,
        auth=(SNOW_USERNAME, SNOW_PASSWORD),
        headers={
            "Accept": "application/json"
        }
    )

    if response.status_code in (200, 204):

        print(
            f"DELETED: {incident['number']} "
            f"- {incident['short_description']}"
        )

        deleted += 1

    else:

        print(
            f"FAILED: {incident['number']} "
            f"- HTTP {response.status_code}"
        )

        print(response.text)

        failed += 1


# --------------------------------------------------
# Summary
# --------------------------------------------------

print()
print("=" * 80)
print("DELETION COMPLETE")
print("=" * 80)

print(f"Total in Excel : {len(incidents)}")
print(f"Deleted        : {deleted}")
print(f"Failed         : {failed}")