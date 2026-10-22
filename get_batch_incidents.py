#This script retrieves incidents from ServiceNow that have an empty assignment group 
# and a short description containing "batch job failure", created between 14 and 28 Sept 2026.
# The results are displayed in the console 
# and exported to an Excel file.


from openpyxl import Workbook
from openpyxl.utils import get_column_letter
import os
import requests
from dotenv import load_dotenv


# --------------------------------------------------
# Load environment variables
# --------------------------------------------------

load_dotenv()

SNOW_URL = os.getenv("SNOW_URL")
SNOW_USERNAME = os.getenv("SNOW_USERNAME")
SNOW_PASSWORD = os.getenv("SNOW_PASSWORD")


# --------------------------------------------------
# ServiceNow URL
# --------------------------------------------------

url = f"{SNOW_URL}/api/now/table/incident"


# --------------------------------------------------
# Export incidents to Excel
# --------------------------------------------------

def export_to_excel(incidents, filename="batch_job_incidents.xlsx"):

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Incidents"

    headers = [
        "Number",
        "Short Description",
        "Assignment Group",
        "State",
        "Created",
        "Sys ID"
    ]

    worksheet.append(headers)

    for incident in incidents:

        worksheet.append([
            incident.get("number", ""),
            incident.get("short_description", ""),
            incident.get("assignment_group", ""),
            incident.get("state", ""),
            incident.get("sys_created_on", ""),
            incident.get("sys_id", "")
        ])

    # Automatically adjust column widths
    for column in worksheet.columns:

        max_length = 0
        column_letter = get_column_letter(column[0].column)

        for cell in column:

            if cell.value:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        worksheet.column_dimensions[column_letter].width = min(
            max_length + 2,
            60
        )

    # Freeze header row
    worksheet.freeze_panes = "A2"

    # Enable filters
    worksheet.auto_filter.ref = worksheet.dimensions

    workbook.save(filename)

    print(f"\nExcel file created: {filename}")


# --------------------------------------------------
# ServiceNow query
# --------------------------------------------------

# Only incidents created between 14 Sept and 28 Sept 2026 (inclusive)
START_DATE = "2026-09-14"
END_DATE = "2026-09-28"

params = {

    "sysparm_query":
        "assignment_groupISEMPTY"
        "^short_descriptionLIKEbatch job failure"
        f"^sys_created_onBETWEENjavascript:gs.dateGenerate('{START_DATE}','00:00:00')"
        f"@javascript:gs.dateGenerate('{END_DATE}','23:59:59')",

    "sysparm_fields":
        "sys_id,number,short_description,assignment_group,state,sys_created_on",

    "sysparm_limit": 1000
}


# --------------------------------------------------
# Get incidents from ServiceNow
# --------------------------------------------------

response = requests.get(
    url,
    auth=(SNOW_USERNAME, SNOW_PASSWORD),
    headers={
        "Accept": "application/json"
    },
    params=params
)


# Raise error if ServiceNow returns an error
response.raise_for_status()


# --------------------------------------------------
# Read response
# --------------------------------------------------

data = response.json()

incidents = data.get("result", [])


# --------------------------------------------------
# Display results
# --------------------------------------------------

print(f"Found {len(incidents)} incidents\n")


for incident in incidents:

    print("=" * 80)

    print(
        f"Number            : "
        f"{incident.get('number')}"
    )

    print(
        f"Short description : "
        f"{incident.get('short_description')}"
    )

    print(
        f"Assignment group  : "
        f"{incident.get('assignment_group')}"
    )

    print(
        f"State             : "
        f"{incident.get('state')}"
    )

    print(
        f"Created           : "
        f"{incident.get('sys_created_on')}"
    )

    print(
        f"Sys ID            : "
        f"{incident.get('sys_id')}"
    )


# --------------------------------------------------
# Export results to Excel
# --------------------------------------------------

export_to_excel(incidents)
